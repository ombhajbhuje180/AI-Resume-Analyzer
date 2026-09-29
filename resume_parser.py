"""
resume_parser.py

Purpose:
    Extract text from uploaded PDF and DOCX resumes.

Supported formats:
    - PDF
    - DOCX

This module does NOT:
    - Train ML models
    - Extract skills
    - Match jobs
    - Generate recommendations
"""

from io import BytesIO

from pypdf import PdfReader
from docx import Document


# ============================================================
# PDF TEXT EXTRACTION
# ============================================================

def extract_pdf_text(file):
    """
    Extract text from a PDF resume.

    Parameters
    ----------
    file:
        Streamlit UploadedFile or file-like object.

    Returns
    -------
    str
        Extracted text from the PDF.
    """

    try:
        # Move file pointer to beginning
        file.seek(0)

        reader = PdfReader(file)

        pages_text = []

        for page in reader.pages:

            text = page.extract_text()

            if text:
                pages_text.append(text)

        return "\n".join(pages_text).strip()

    except Exception as error:

        raise RuntimeError(
            f"Failed to read PDF file: {error}"
        )


# ============================================================
# DOCX TEXT EXTRACTION
# ============================================================

def extract_docx_text(file):
    """
    Extract text from a DOCX resume.

    Parameters
    ----------
    file:
        Streamlit UploadedFile or file-like object.

    Returns
    -------
    str
        Extracted text from the DOCX.
    """

    try:
        # Read uploaded file into memory
        file.seek(0)

        file_bytes = file.read()

        document = Document(
            BytesIO(file_bytes)
        )

        paragraphs = []

        # ----------------------------------------------------
        # Extract normal paragraphs
        # ----------------------------------------------------

        for paragraph in document.paragraphs:

            text = paragraph.text.strip()

            if text:
                paragraphs.append(text)

        # ----------------------------------------------------
        # Extract text from tables
        # ----------------------------------------------------

        for table in document.tables:

            for row in table.rows:

                row_text = []

                for cell in row.cells:

                    cell_text = cell.text.strip()

                    if cell_text:
                        row_text.append(cell_text)

                if row_text:

                    paragraphs.append(
                        " | ".join(row_text)
                    )

        return "\n".join(paragraphs).strip()

    except Exception as error:

        raise RuntimeError(
            f"Failed to read DOCX file: {error}"
        )


# ============================================================
# GENERAL RESUME EXTRACTION
# ============================================================

def extract_resume_text(file, file_extension=None):
    """
    Extract text from a PDF or DOCX resume.

    Parameters
    ----------
    file:
        Streamlit UploadedFile or file-like object.

    file_extension:
        File extension such as 'pdf' or 'docx'.

    Returns
    -------
    str
        Extracted resume text.
    """

    # --------------------------------------------------------
    # Get extension automatically if not provided
    # --------------------------------------------------------

    if file_extension is None:

        if hasattr(file, "name"):

            file_extension = (
                file.name
                .split(".")[-1]
                .lower()
            )

        else:

            raise ValueError(
                "Unable to determine file type."
            )


    file_extension = (
        file_extension
        .lower()
        .replace(".", "")
    )


    # --------------------------------------------------------
    # PDF
    # --------------------------------------------------------

    if file_extension == "pdf":

        return extract_pdf_text(file)


    # --------------------------------------------------------
    # DOCX
    # --------------------------------------------------------

    elif file_extension == "docx":

        return extract_docx_text(file)


    # --------------------------------------------------------
    # Unsupported format
    # --------------------------------------------------------

    else:

        raise ValueError(
            f"Unsupported file format: "
            f"{file_extension}. "
            f"Only PDF and DOCX are supported."
        )


# ============================================================
# BASIC TEXT VALIDATION
# ============================================================

def validate_resume_text(text):
    """
    Perform basic validation on extracted resume text.

    Returns True when the extracted text appears usable.
    """

    if not text:

        return False

    cleaned_text = text.strip()

    # Very small amount of text probably means
    # extraction failed.

    if len(cleaned_text) < 50:

        return False

    return True


# ============================================================
# RESUME INFORMATION
# ============================================================

def get_resume_statistics(text):
    """
    Calculate basic statistics from extracted resume text.

    Returns
    -------
    dict
        Basic resume statistics.
    """

    if not text:

        return {
            "characters": 0,
            "words": 0,
            "lines": 0
        }


    words = text.split()

    lines = [
        line
        for line in text.splitlines()
        if line.strip()
    ]


    return {
        "characters": len(text),
        "words": len(words),
        "lines": len(lines)
    }


# ============================================================
# TESTING
# ============================================================

if __name__ == "__main__":

    print("=" * 60)

    print(
        "RESUME PARSER MODULE"
    )

    print("=" * 60)

    print(
        "\nThis module is designed to be imported by app.py."
    )

    print(
        "\nSupported formats:"
    )

    print(
        "  ✓ PDF"
    )

    print(
        "  ✓ DOCX"
    )

    print(
        "\nMain function:"
    )

    print(
        "  extract_resume_text(file)"
    )

    print(
        "\nResume parser ready ✅"
    )

    print("=" * 60)